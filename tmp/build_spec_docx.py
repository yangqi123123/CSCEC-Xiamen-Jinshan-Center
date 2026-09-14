import json
import re
import subprocess
from copy import deepcopy
from pathlib import Path

from lxml import html
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\10208\Desktop\四局厦门金山项目原型")
TEMPLATE = ROOT / "docs" / "厦门金山财富中心需求规格说明书V1.0 .docx"
SOURCE_MD = ROOT / "docs" / "产品概要设计文档-厦门金山财富中心.md"
OUTPUT = ROOT / "docs" / "厦门金山财富中心需求规格说明书V1.0-完整版.docx"
NODE = Path(r"C:\Users\10208\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, rows, widths=None):
    if not rows:
        return
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    for ri, values in enumerate(rows):
        cells = table.add_row().cells
        for ci, value in enumerate(values):
            cells[ci].text = str(value)
            if widths:
                set_cell_width(cells[ci], widths[ci])
            for paragraph in cells[ci].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                    run.font.size = Pt(9)
                    if ri == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        if ri == 0:
            for cell in cells:
                shade(cell, "2F75B5")
            set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {min(level, 3)}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)
    return p


def add_bullet(doc, text, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(3)
    return p


def parse_markdown_until_details(doc):
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines = []
    details_added = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph("\n".join(code_lines), style="HTML Preformatted")
                p.paragraph_format.space_after = Pt(6)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                vals = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", x) for x in vals):
                    block.append(vals)
                i += 1
            add_table(doc, block)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            text = m.group(2)
            if text == "厦门金山财富中心 IoT 管理台":
                i += 1
                continue
            level = len(m.group(1))
            if level == 2 and text == "产品概要设计文档":
                i += 1
                continue
            if level == 2 and text == "六、产品迭代计划" and not details_added:
                add_big_screen_details(doc)
                add_backend_details(doc)
                details_added = True
            add_heading(doc, text, max(1, level - 1))
        elif re.match(r"^\d+\.\s+", line):
            add_bullet(doc, re.sub(r"^\d+\.\s+", "", line), True)
        elif line.startswith("+ "):
            add_bullet(doc, line[2:])
        elif line.strip() and line.strip() != "---":
            add_body(doc, re.sub(r"\*\*", "", line).replace("  ", " "))
        i += 1


def backend_docs():
    source = (ROOT / "app" / "components" / "requirement.js").read_text(encoding="utf-8")
    start = source.index("(function () {") + len("(function () {")
    end = source.index("  const escapeHtml")
    code = source[start:end] + "\nconsole.log(JSON.stringify(docs));"
    result = subprocess.run([str(NODE), "-e", code], capture_output=True, text=True, encoding="utf-8", check=True)
    return json.loads(result.stdout)


def add_backend_details(doc):
    docs = backend_docs()
    add_heading(doc, "5.4 后台页面功能详细规范", 2)
    add_body(doc, "以下内容直接取自后台 HTML 原型的页面需求说明配置。未在页面中定义的功能不作扩展。")
    for index, item in enumerate(docs.values(), 1):
        add_heading(doc, f"5.4.{index} {item['title']}", 3)
        for title, content in item["sections"]:
            if content is None or content == "":
                content = "无"
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = title == "字段说明"
            run = p.add_run(title)
            run.bold = True
            run.font.name = "黑体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            if title == "字段说明":
                rows = [["字段名称", "类型", "是否必填", "默认值", "取值范围"]] + content
                add_table(doc, rows, [1800, 1050, 1050, 1100, 3800])
            else:
                add_body(doc, str(content))


def add_big_screen_details(doc):
    add_heading(doc, "5.3 大屏页面功能详细规范", 2)
    page_files = [
        ("总览", "overview.html"), ("机电系统", "mep.html"), ("弱电系统", "weak-electric.html"),
        ("能源管理", "energy.html"), ("运营管理", "operation.html")
    ]
    for idx, (name, filename) in enumerate(page_files, 1):
        tree = html.fromstring((ROOT / "Big Screen" / filename).read_text(encoding="utf-8"))
        matches = tree.xpath('//template[@id="requirementDocument"]')
        template = matches[0] if matches else None
        add_heading(doc, f"5.3.{idx} {name}", 3)
        if not template:
            add_body(doc, "【待补充】")
            continue
        sections = template.xpath('.//*[@data-requirement-section]')
        for section in sections:
            headings = section.xpath('.//h2|.//h3|.//h4')
            heading = headings[0] if headings else None
            title = " ".join(heading.itertext()).strip() if heading is not None else "需求内容"
            add_heading(doc, title, 3)
            for table in section.xpath('.//table'):
                rows = [[" ".join(cell.itertext()).strip() for cell in row.xpath('./th|./td')] for row in table.xpath('.//tr')]
                table.getparent().remove(table)
                add_table(doc, rows)
            text = "\n".join(x.strip() for x in section.itertext() if x.strip())
            if heading is not None:
                text = text.replace(title, "", 1).strip()
            if text:
                for part in text.split("\n"):
                    if part.strip():
                        add_body(doc, part.strip())
            elif not section.xpath('.//table'):
                add_body(doc, "【待补充】")


def remove_placeholder_content(doc):
    body = doc._element.body
    marker = None
    for p in doc.paragraphs:
        if p.text.strip() == "项目概述":
            marker = p._element
            break
    if marker is None:
        return
    found = False
    for child in list(body):
        if child is marker:
            found = True
        if found and child.tag != qn("w:sectPr"):
            body.remove(child)


def style_document(doc):
    section = doc.sections[-1]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for level, size in [(1, 16), (2, 14), (3, 12)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    for p in doc.paragraphs:
        for r in p.runs:
            if not r.font.name:
                r.font.name = "宋体"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def main():
    doc = Document(TEMPLATE)
    remove_placeholder_content(doc)
    style_document(doc)
    # Complete front matter in the supplied template.
    doc.tables[0].cell(0, 2).text = "XMJS-IOT-SRS-001"
    doc.tables[0].cell(2, 2).text = "2026-09-14"
    doc.tables[2].cell(1, 1).text = "2026-09-14"
    parse_markdown_until_details(doc)
    style_document(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
